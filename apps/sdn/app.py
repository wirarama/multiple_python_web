"""
Simulator Node Jaringan & Evaluasi Algoritma Optimasi Routing SDN.

Flask + NetworkX. Menyediakan:
- Editor topologi (tambah/hapus node & link, atur delay/bandwidth/loss tiap link)
- Formulir parameter khusus tiap algoritma (dijalankan dari UI, dikirim via API)
- Evaluasi metrik: Delay, Hop Count, Throughput (simulasi), Packet Loss (simulasi)
- Klasifikasi & skor QoS (qos.py)
- Visualisasi topologi + jalur hasil optimasi

10 algoritma: Dijkstra, A*, ACO, Genetic Algorithm, PSO, Simulated Annealing,
Tabu Search, Q-Learning, Greedy Routing, dan Minimum Spanning Tree.

Jalankan:
    pip install -r requirements.txt
    python app.py
Lalu buka http://127.0.0.1:5000
"""
import os
import json
import random
from io import BytesIO
from datetime import datetime
import networkx as nx
from flask import Flask, jsonify, request, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from graph_utils import load_topology, graph_to_dict, DEFAULT_TOPOLOGY_PATH
from metrics import compute_path_metrics, compute_mst_metrics
from algorithms import ALGORITHMS, clean_params
from topology_generator import generate_random_topology
import qos as qos_module
from algorithms.load_balancing import (
    LB_ALGORITHMS, generate_cluster, evaluate_assignment, reference as lb_reference,
    classify as lb_classify, clean_params as lb_clean_params, client_log as lb_client_log,
    recompute_latency as lb_recompute_latency, scaling_suggestions as lb_scaling_suggestions,
)
import intrusion_detection as ids_module

app = Flask(__name__)

G, RAW_TOPOLOGY = load_topology()
LAYOUT = nx.spring_layout(G, seed=42)
nx.set_node_attributes(G, LAYOUT, "pos")

# Mode topologi aktif: "default" (dari data/topology.json atau diedit manual,
# posisi dihitung via spring_layout) atau "generated" (hasil simulasi jumlah
# node + luas wilayah, posisi = koordinat meter sebenarnya pada area tsb).
TOPO_MODE = "default"
AREA = {"width": 1000.0, "height": 1000.0}

# Cluster load balancing aktif (server, klien, request, latensi) -- lihat
# algorithms/load_balancing.py. Dibangkitkan on-demand lewat /api/loadbalance/generate.
LB_CLUSTER = generate_cluster(seed=42)

# Skenario simulasi deteksi intrusi aktif (server, client umum/penyerang) --
# lihat intrusion_detection.py. Dibangkitkan on-demand lewat /api/ids/generate.
IDS_CLUSTER = ids_module.generate_scenario(seed=42)


def _full_relayout():
    """Hitung ulang posisi SELURUH node dari nol (spring layout, mode
    'default'). Dipakai hanya saat topologi benar-benar diganti (reset,
    generate, replace) -- BUKAN saat sekadar tambah/hapus node/link, supaya
    posisi node yang sudah digeser manual oleh pengguna (drag & drop) tidak
    hilang begitu saja."""
    global LAYOUT
    LAYOUT = nx.spring_layout(G, seed=42)
    nx.set_node_attributes(G, LAYOUT, "pos")


def _ensure_positions():
    """Pastikan seluruh node punya atribut 'pos', TANPA mengubah posisi node
    yang sudah punya (termasuk yang sudah digeser manual via drag & drop pada
    visualisasi). Node baru diberi koordinat acak: dalam AREA aktif pada mode
    'generated', atau di sekitar bounding box layout yang sudah ada pada mode
    'default'. Dipanggil setelah tambah/hapus node/link agar layout lama
    dipertahankan (relayout penuh hanya terjadi lewat _full_relayout)."""
    global LAYOUT
    rng = random.Random()
    pos = nx.get_node_attributes(G, "pos")
    if TOPO_MODE == "generated":
        x_lo, x_hi = 0.0, AREA["width"]
        y_lo, y_hi = 0.0, AREA["height"]
    elif pos:
        xs = [p[0] for p in pos.values()]
        ys = [p[1] for p in pos.values()]
        x_lo, x_hi = min(xs), max(xs)
        y_lo, y_hi = min(ys), max(ys)
    else:
        x_lo, x_hi, y_lo, y_hi = -1.0, 1.0, -1.0, 1.0
    for n in G.nodes():
        if n not in pos:
            pos[n] = (rng.uniform(x_lo, x_hi), rng.uniform(y_lo, y_hi))
    pos = {n: pos[n] for n in G.nodes()}   # buang node yang sudah dihapus
    nx.set_node_attributes(G, pos, "pos")
    LAYOUT = pos


def _run_single(algo_key, source, target, raw_params, seed, congestion):
    if algo_key not in ALGORITHMS:
        return None, {"error": f"Algoritma '{algo_key}' tidak dikenal"}

    spec = ALGORITHMS[algo_key]
    module = spec["module"]
    params = clean_params(algo_key, raw_params or {})
    ref = qos_module.topology_reference(G)

    if spec["type"] == "tree":
        result = module.run(G, **params)
        metrics = compute_mst_metrics(G, result["mst_edges"], seed=seed, congestion=congestion)
        payload = {
            "algorithm": algo_key,
            "label": spec["label"],
            "type": "tree",
            "cost": result.get("cost"),
            "runtime_ms": result.get("runtime_ms"),
            "meta": result.get("meta", {}),
            "params": params,
            "metrics": metrics,
            "qos": qos_module.classify(metrics, ref),
        }
        return payload, None

    if source not in G or target not in G:
        return None, {"error": "Node source/target tidak ditemukan pada topologi"}

    result = module.run(G, source, target, seed=seed, **params)
    path = result.get("path")
    metrics = compute_path_metrics(G, path, seed=seed, congestion=congestion)
    payload = {
        "algorithm": algo_key,
        "label": spec["label"],
        "type": "path",
        "path": path,
        "cost": result.get("cost"),
        "runtime_ms": result.get("runtime_ms"),
        "meta": result.get("meta", {}),
        "params": params,
        "metrics": metrics,
        "qos": qos_module.classify(metrics, ref),
    }
    return payload, None


@app.route("/")
def index():
    return render_template(
        "index.html",
        nodes=sorted(G.nodes()),
        algorithms=[
            {"key": k, "label": v["label"], "type": v["type"], "params": v["params"]}
            for k, v in ALGORITHMS.items()
        ],
    )


@app.route("/tutorial")
def tutorial():
    """Halaman panduan pengguna (tutorial UI, simulasi & algoritma), dibuka
    lewat tautan "?" pada tiap komponen UI atau langsung via /tutorial."""
    return render_template("tutorial.html")


# ---------------------------------------------------------------- topologi --

@app.route("/api/topology")
def api_topology():
    resp = graph_to_dict(G, layout=LAYOUT)
    resp["mode"] = TOPO_MODE
    resp["area"] = AREA
    return jsonify(resp)


@app.route("/api/topology", methods=["POST"])
def api_topology_replace():
    """Ganti seluruh topologi aktif dengan JSON yang dikirim pada body request."""
    global G, RAW_TOPOLOGY, TOPO_MODE
    data = request.get_json(force=True)
    if not data or "nodes" not in data or "edges" not in data:
        return jsonify({"error": "Format topologi tidak valid (butuh 'nodes' dan 'edges')"}), 400

    tmp_path = os.path.join(os.path.dirname(__file__), "data", "_uploaded_topology.json")
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(data, f)

    G, RAW_TOPOLOGY = load_topology(tmp_path)
    TOPO_MODE = "default"
    _full_relayout()
    return jsonify(graph_to_dict(G, layout=LAYOUT))


@app.route("/api/topology/reset", methods=["POST"])
def api_topology_reset():
    """Muat ulang topologi bawaan (data/topology.json)."""
    global G, RAW_TOPOLOGY, TOPO_MODE
    G, RAW_TOPOLOGY = load_topology(DEFAULT_TOPOLOGY_PATH)
    TOPO_MODE = "default"
    _full_relayout()
    return jsonify(graph_to_dict(G, layout=LAYOUT))


@app.route("/api/topology/generate", methods=["POST"])
def api_topology_generate():
    """Bangkitkan topologi acak berdasarkan parameter simulasi: jumlah node
    dan luas wilayah (lebar x tinggi, meter). Menggantikan topologi aktif."""
    global G, RAW_TOPOLOGY, TOPO_MODE, AREA, LAYOUT
    data = request.get_json(force=True) or {}
    try:
        num_nodes = int(data.get("num_nodes", 10))
        area_width = float(data.get("area_width", 1000))
        area_height = float(data.get("area_height", 1000))
        k_neighbors = int(data.get("k_neighbors", 3))
    except (TypeError, ValueError):
        return jsonify({"error": "num_nodes/area_width/area_height/k_neighbors harus berupa angka"}), 400
    seed = data.get("seed")

    if not (2 <= num_nodes <= 300):
        return jsonify({"error": "num_nodes harus di rentang 2-300"}), 400
    if area_width <= 0 or area_height <= 0:
        return jsonify({"error": "area_width dan area_height harus > 0"}), 400

    G, pos = generate_random_topology(num_nodes, area_width, area_height, k_neighbors, seed)
    nx.set_node_attributes(G, pos, "pos")
    LAYOUT = pos
    TOPO_MODE = "generated"
    AREA = {"width": area_width, "height": area_height}
    RAW_TOPOLOGY = {"nodes": list(G.nodes()), "edges": []}

    resp = graph_to_dict(G, layout=LAYOUT)
    resp["mode"] = TOPO_MODE
    resp["area"] = AREA
    return jsonify(resp)


@app.route("/api/topology/status")
def api_topology_status():
    return jsonify({
        "mode": TOPO_MODE, "area": AREA,
        "num_nodes": G.number_of_nodes(), "num_edges": G.number_of_edges(),
    })


@app.route("/api/topology/node", methods=["POST"])
def api_add_node():
    data = request.get_json(force=True) or {}
    node_id = str(data.get("id", "")).strip()
    if not node_id:
        return jsonify({"error": "id node tidak boleh kosong"}), 400
    if node_id in G:
        return jsonify({"error": f"Node '{node_id}' sudah ada"}), 400
    G.add_node(node_id)
    _ensure_positions()
    return jsonify(graph_to_dict(G, layout=LAYOUT))


@app.route("/api/topology/node/<node_id>", methods=["DELETE"])
def api_delete_node(node_id):
    if node_id not in G:
        return jsonify({"error": f"Node '{node_id}' tidak ditemukan"}), 404
    if G.number_of_nodes() <= 2:
        return jsonify({"error": "Topologi minimal harus memiliki 2 node"}), 400
    G.remove_node(node_id)
    _ensure_positions()
    return jsonify(graph_to_dict(G, layout=LAYOUT))


@app.route("/api/topology/edge", methods=["POST"])
def api_upsert_edge():
    data = request.get_json(force=True) or {}
    u, v = data.get("source"), data.get("target")
    if not u or not v:
        return jsonify({"error": "source dan target wajib diisi"}), 400
    if u == v:
        return jsonify({"error": "source dan target tidak boleh sama (self-loop)"}), 400
    for n in (u, v):
        if n not in G:
            G.add_node(n)
    try:
        delay = float(data.get("delay", 1.0))
        bandwidth = float(data.get("bandwidth", 100.0))
        loss = float(data.get("loss", 0.001))
    except (TypeError, ValueError):
        return jsonify({"error": "delay/bandwidth/loss harus berupa angka"}), 400
    if delay <= 0 or bandwidth <= 0 or not (0 <= loss < 1):
        return jsonify({"error": "delay & bandwidth harus > 0, loss harus di rentang [0,1)"}), 400

    G.add_edge(u, v, delay=delay, bandwidth=bandwidth, loss=loss, weight=delay)
    _ensure_positions()
    return jsonify(graph_to_dict(G, layout=LAYOUT))


@app.route("/api/topology/edge", methods=["DELETE"])
def api_delete_edge():
    data = request.get_json(force=True) or {}
    u, v = data.get("source"), data.get("target")
    if not G.has_edge(u, v):
        return jsonify({"error": f"Link {u}-{v} tidak ditemukan"}), 404
    G.remove_edge(u, v)
    _ensure_positions()
    return jsonify(graph_to_dict(G, layout=LAYOUT))


@app.route("/api/topology/positions", methods=["POST"])
def api_topology_positions():
    """Perbarui posisi (x, y) satu atau beberapa node hasil digeser (drag &
    drop) pada visualisasi topologi. Posisi baru langsung dipakai ulang oleh
    heuristik A* (jarak Euclidean antar node) pada evaluasi berikutnya, dan
    ikut tersimpan saat konfigurasi diekspor ke JSON."""
    data = request.get_json(force=True) or {}
    positions = data.get("positions") or {}
    if not isinstance(positions, dict) or not positions:
        return jsonify({"error": "positions wajib berupa objek {node_id: [x, y]}"}), 400

    pos = nx.get_node_attributes(G, "pos")
    updated = 0
    for node_id, xy in positions.items():
        if node_id not in G:
            continue
        try:
            x, y = float(xy[0]), float(xy[1])
        except (TypeError, ValueError, IndexError, KeyError):
            return jsonify({"error": f"Posisi node '{node_id}' tidak valid"}), 400
        pos[node_id] = (x, y)
        updated += 1

    if not updated:
        return jsonify({"error": "Tidak ada node valid pada payload positions"}), 400

    global LAYOUT
    nx.set_node_attributes(G, pos, "pos")
    LAYOUT = pos
    return jsonify(graph_to_dict(G, layout=LAYOUT))


# --------------------------------------------------- konfigurasi (export/import) --

@app.route("/api/config/export")
def api_config_export():
    """Ekspor seluruh konfigurasi aktif (topologi routing + cluster load
    balancing) sebagai satu paket JSON yang bisa diimpor kembali nanti."""
    return jsonify({
        "version": 1,
        "exported_at": datetime.now().isoformat(),
        "topology": {
            "mode": TOPO_MODE,
            "area": AREA,
            "graph": graph_to_dict(G, layout=LAYOUT),
        },
        "load_balancing": LB_CLUSTER,
    })


@app.route("/api/config/import", methods=["POST"])
def api_config_import():
    """Impor konfigurasi (hasil /api/config/export) dan terapkan sebagai
    topologi + cluster load balancing aktif. Bagian 'topology' dan
    'load_balancing' opsional -- bila salah satu tidak ada di body, bagian
    tersebut tidak diubah."""
    global G, RAW_TOPOLOGY, TOPO_MODE, AREA, LAYOUT, LB_CLUSTER
    data = request.get_json(force=True) or {}
    topo = data.get("topology")
    lb = data.get("load_balancing")

    if topo:
        graph = topo.get("graph") or {}
        nodes = graph.get("nodes")
        edges = graph.get("edges")
        if not nodes or edges is None:
            return jsonify({"error": "Konfigurasi topologi tidak valid (butuh 'graph.nodes' & 'graph.edges')"}), 400
        try:
            raw = {
                "nodes": [n["id"] for n in nodes],
                "edges": [
                    {
                        "source": e["source"], "target": e["target"],
                        "delay": float(e.get("delay", 1.0)),
                        "bandwidth": float(e.get("bandwidth", 100.0)),
                        "loss": float(e.get("loss", 0.001)),
                    }
                    for e in edges
                ],
            }
        except (KeyError, TypeError, ValueError):
            return jsonify({"error": "Konfigurasi topologi tidak valid"}), 400

        tmp_path = os.path.join(os.path.dirname(__file__), "data", "_uploaded_topology.json")
        with open(tmp_path, "w", encoding="utf-8") as f:
            json.dump(raw, f)
        G, RAW_TOPOLOGY = load_topology(tmp_path)
        TOPO_MODE = topo.get("mode", "default")
        AREA = topo.get("area", AREA)
        # Posisi (x, y) tiap node -- termasuk hasil digeser manual (drag & drop)
        # sebelum diekspor -- ikut dipulihkan untuk kedua mode topologi.
        try:
            pos = {
                n["id"]: (float(n["x"]), float(n["y"]))
                for n in nodes if "x" in n and "y" in n
            }
        except (KeyError, TypeError, ValueError):
            pos = {}
        nx.set_node_attributes(G, pos, "pos")
        _ensure_positions()

    if lb:
        required = {"servers", "clients", "requests", "latency", "area"}
        if not required.issubset(lb.keys()):
            return jsonify({"error": "Konfigurasi load balancing tidak valid (butuh servers/clients/requests/latency/area)"}), 400
        LB_CLUSTER = lb

    return jsonify({
        "topology": graph_to_dict(G, layout=LAYOUT),
        "topology_mode": TOPO_MODE,
        "area": AREA,
        "load_balancing": _lb_topology_dict(),
    })


# --------------------------------------------------------------- algoritma --

@app.route("/api/algorithms")
def api_algorithms():
    return jsonify([
        {"key": k, "label": v["label"], "type": v["type"], "params": v["params"]}
        for k, v in ALGORITHMS.items()
    ])


@app.route("/api/evaluate", methods=["POST"])
def api_evaluate():
    data = request.get_json(force=True) or {}
    algo_key = data.get("algorithm")
    source = data.get("source")
    target = data.get("target")
    params = data.get("params", {}) or {}
    seed = data.get("seed", 42)
    congestion = float(data.get("congestion", 0.0))

    payload, error = _run_single(algo_key, source, target, params, seed, congestion)
    if error:
        return jsonify(error), 400
    return jsonify(payload)


@app.route("/api/compare", methods=["POST"])
def api_compare():
    data = request.get_json(force=True) or {}
    source = data.get("source")
    target = data.get("target")
    seed = data.get("seed", 42)
    congestion = float(data.get("congestion", 0.0))
    selected = data.get("algorithms") or list(ALGORITHMS.keys())
    per_algo_params = data.get("params", {}) or {}   # {"aco": {...}, "genetic": {...}, ...}

    results = []
    for algo_key in selected:
        payload, error = _run_single(
            algo_key, source, target, per_algo_params.get(algo_key, {}), seed, congestion
        )
        if error:
            results.append({"algorithm": algo_key, "label": ALGORITHMS.get(algo_key, {}).get("label", algo_key), "error": error.get("error")})
        else:
            results.append(payload)
    return jsonify({"source": source, "target": target, "congestion": congestion, "results": results})


# ---------------------------------------------------------- load balancing --

def _lb_topology_dict():
    c = LB_CLUSTER
    return {
        "servers": c["servers"],
        "clients": c["clients"],
        "requests": c["requests"],
        "num_requests": len(c["requests"]),
        "area": c["area"],
    }


def _run_single_lb(algo_key, raw_params, seed):
    if algo_key not in LB_ALGORITHMS:
        return None, {"error": f"Algoritma load balancing '{algo_key}' tidak dikenal"}
    spec = LB_ALGORITHMS[algo_key]
    params = lb_clean_params(algo_key, raw_params or {})
    result = spec["run"](LB_CLUSTER, seed=seed, **params)
    metrics = evaluate_assignment(LB_CLUSTER, result.get("assignment"))
    ref = lb_reference(LB_CLUSTER)
    request_log = metrics.pop("request_log", None)
    payload = {
        "algorithm": algo_key,
        "label": spec["label"],
        "cost": result.get("cost"),
        "runtime_ms": result.get("runtime_ms"),
        "meta": result.get("meta", {}),
        "params": params,
        "metrics": metrics,
        "qos": lb_classify(metrics, ref),
        "client_log": lb_client_log(LB_CLUSTER, request_log) if metrics.get("valid") else [],
        "scaling_suggestions": lb_scaling_suggestions(LB_CLUSTER, metrics),
    }
    return payload, None


@app.route("/api/loadbalance/topology")
def api_lb_topology():
    return jsonify(_lb_topology_dict())


@app.route("/api/loadbalance/positions", methods=["POST"])
def api_lb_positions():
    """Perbarui posisi (x, y) satu/beberapa server dan/atau klien hasil
    digeser (drag & drop) pada visualisasi cluster. Latensi jaringan
    klien->server (dipakai algoritma load balancing utk komponen response
    time) dihitung ulang dari posisi baru, dan posisi ikut tersimpan saat
    konfigurasi diekspor ke JSON."""
    data = request.get_json(force=True) or {}
    server_pos = data.get("servers") or {}
    client_pos = data.get("clients") or {}
    if not isinstance(server_pos, dict) or not isinstance(client_pos, dict):
        return jsonify({"error": "servers/clients wajib berupa objek {id: [x, y]}"}), 400
    if not server_pos and not client_pos:
        return jsonify({"error": "Tidak ada posisi yang dikirim"}), 400

    updated = 0
    for entry_list, pos_map in ((LB_CLUSTER["servers"], server_pos), (LB_CLUSTER["clients"], client_pos)):
        by_id = {e["id"]: e for e in entry_list}
        for node_id, xy in pos_map.items():
            if node_id not in by_id:
                continue
            try:
                x, y = float(xy[0]), float(xy[1])
            except (TypeError, ValueError, IndexError, KeyError):
                return jsonify({"error": f"Posisi '{node_id}' tidak valid"}), 400
            by_id[node_id]["x"] = x
            by_id[node_id]["y"] = y
            updated += 1

    if not updated:
        return jsonify({"error": "Tidak ada server/klien valid pada payload"}), 400

    lb_recompute_latency(LB_CLUSTER)
    return jsonify(_lb_topology_dict())


@app.route("/api/loadbalance/generate", methods=["POST"])
def api_lb_generate():
    global LB_CLUSTER
    data = request.get_json(force=True) or {}
    try:
        num_servers = int(data.get("num_servers", 6))
        num_clients = int(data.get("num_clients", 8))
        requests_per_client = int(data.get("requests_per_client", 4))
        area_width = float(data.get("area_width", 1000))
        area_height = float(data.get("area_height", 1000))
        capacity_min = float(data.get("capacity_min", 60))
        capacity_max = float(data.get("capacity_max", 220))
        demand_min = float(data.get("demand_min", 1))
        demand_max = float(data.get("demand_max", 20))
    except (TypeError, ValueError):
        return jsonify({"error": "Semua parameter simulasi cluster harus berupa angka"}), 400
    seed = data.get("seed")

    if not (1 <= num_servers <= 100):
        return jsonify({"error": "num_servers harus di rentang 1-100"}), 400
    if not (1 <= num_clients <= 200):
        return jsonify({"error": "num_clients harus di rentang 1-200"}), 400
    if not (1 <= requests_per_client <= 200):
        return jsonify({"error": "requests_per_client harus di rentang 1-200"}), 400
    if area_width <= 0 or area_height <= 0:
        return jsonify({"error": "area_width dan area_height harus > 0"}), 400
    if capacity_min <= 0 or capacity_max < capacity_min:
        return jsonify({"error": "rentang capacity tidak valid"}), 400
    if demand_min <= 0 or demand_max < demand_min:
        return jsonify({"error": "rentang demand tidak valid"}), 400

    LB_CLUSTER = generate_cluster(
        num_servers=num_servers, num_clients=num_clients, requests_per_client=requests_per_client,
        area_width=area_width, area_height=area_height,
        capacity_range=(capacity_min, capacity_max), demand_range=(demand_min, demand_max),
        seed=seed,
    )
    return jsonify(_lb_topology_dict())


@app.route("/api/loadbalance/algorithms")
def api_lb_algorithms():
    return jsonify([
        {"key": k, "label": v["label"], "params": v["params"]}
        for k, v in LB_ALGORITHMS.items()
    ])


@app.route("/api/loadbalance/evaluate", methods=["POST"])
def api_lb_evaluate():
    data = request.get_json(force=True) or {}
    algo_key = data.get("algorithm")
    params = data.get("params", {}) or {}
    seed = data.get("seed", 42)
    payload, error = _run_single_lb(algo_key, params, seed)
    if error:
        return jsonify(error), 400
    return jsonify(payload)


@app.route("/api/loadbalance/compare", methods=["POST"])
def api_lb_compare():
    data = request.get_json(force=True) or {}
    seed = data.get("seed", 42)
    selected = data.get("algorithms") or list(LB_ALGORITHMS.keys())
    per_algo_params = data.get("params", {}) or {}

    results = []
    for algo_key in selected:
        payload, error = _run_single_lb(algo_key, per_algo_params.get(algo_key, {}), seed)
        if error:
            results.append({"algorithm": algo_key, "label": LB_ALGORITHMS.get(algo_key, {}).get("label", algo_key), "error": error.get("error")})
        else:
            results.append(payload)
    return jsonify({"seed": seed, "results": results})


# ------------------------------------------------------- intrusion detection --

def _ids_topology_dict():
    c = IDS_CLUSTER
    return {
        "servers": c["servers"],
        "clients": c["clients"],
        "area": c["area"],
        "attacker_ratio": c["attacker_ratio"],
        "num_attackers": c["num_attackers"],
        "num_normal": c["num_normal"],
    }


@app.route("/api/ids/attack-types")
def api_ids_attack_types():
    """Daftar jenis intrusion yang dapat dicentang di UI (DoS/DDoS, Probing,
    R2L, U2R) beserta deskripsi & potensi/intensitas default."""
    return jsonify(ids_module.attack_type_list())


@app.route("/api/ids/topology")
def api_ids_topology():
    return jsonify(_ids_topology_dict())


@app.route("/api/ids/generate", methods=["POST"])
def api_ids_generate():
    """Bangkitkan topologi server & client (umum/penyerang) baru untuk
    simulasi deteksi intrusi, berdasarkan jumlah server, jumlah client, dan
    proporsi client penyerang (attacker_ratio, 0-1)."""
    global IDS_CLUSTER
    data = request.get_json(force=True) or {}
    try:
        num_servers = int(data.get("num_servers", 4))
        num_clients = int(data.get("num_clients", 20))
        attacker_ratio = float(data.get("attacker_ratio", 0.25))
        area_width = float(data.get("area_width", 1000))
        area_height = float(data.get("area_height", 1000))
    except (TypeError, ValueError):
        return jsonify({"error": "num_servers/num_clients/attacker_ratio/area_width/area_height harus berupa angka"}), 400
    seed = data.get("seed")

    if not (1 <= num_servers <= 30):
        return jsonify({"error": "num_servers harus di rentang 1-30"}), 400
    if not (1 <= num_clients <= 300):
        return jsonify({"error": "num_clients harus di rentang 1-300"}), 400
    if not (0 <= attacker_ratio <= 1):
        return jsonify({"error": "attacker_ratio harus di rentang 0-1"}), 400
    if area_width <= 0 or area_height <= 0:
        return jsonify({"error": "area_width dan area_height harus > 0"}), 400

    IDS_CLUSTER = ids_module.generate_scenario(
        num_servers=num_servers, num_clients=num_clients, attacker_ratio=attacker_ratio,
        area_width=area_width, area_height=area_height, seed=seed,
    )
    return jsonify(_ids_topology_dict())


@app.route("/api/ids/simulate", methods=["POST"])
def api_ids_simulate():
    """Bangkitkan traffic/flow acak client->server pada topologi IDS aktif dan
    jalankan detector berbasis aturan sederhana. `attacks` = daftar key jenis
    intrusion yang dicentang (dos_ddos/probing/r2l/u2r); `potensi` = dict
    key->0..1 (intensitas tiap jenis, dari slider/persentase di UI). Bila
    `attacks` kosong, seluruh traffic (termasuk dari client penyerang)
    dianggap normal (tidak ada serangan yang disimulasikan)."""
    data = request.get_json(force=True) or {}
    selected = data.get("attacks") or []
    potensi = data.get("potensi") or {}
    try:
        flows_per_client = int(data.get("flows_per_client", 6))
    except (TypeError, ValueError):
        return jsonify({"error": "flows_per_client harus berupa angka"}), 400
    if not (1 <= flows_per_client <= 50):
        return jsonify({"error": "flows_per_client harus di rentang 1-50"}), 400
    seed = data.get("seed", 42)

    flows, summary = ids_module.simulate_traffic(
        IDS_CLUSTER, selected, potensi, flows_per_client=flows_per_client, seed=seed
    )
    return jsonify({"flows": flows, "summary": summary, "attacks": selected, "potensi": potensi})


# ------------------------------------------------------------- export excel --

def _autosize_columns(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(60, max(10, length + 2))


def _write_header(ws, headers):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2A78D6")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


@app.route("/api/export/excel", methods=["POST"])
def api_export_excel():
    """Ekspor hasil evaluasi (routing dan/atau load balancing) yang sedang
    ditampilkan di UI menjadi satu file Excel (.xlsx) multi-sheet. Data
    hasil run tidak disimpan di backend, sehingga dikirim dari frontend
    (hasil run terakhir + info topologi/cluster terkait) via body request."""
    data = request.get_json(force=True) or {}
    routing = data.get("routing") or {}
    lb = data.get("load_balancing") or {}

    if not routing.get("results") and not lb.get("results"):
        return jsonify({"error": "Tidak ada hasil untuk diekspor. Jalankan perbandingan terlebih dahulu."}), 400

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Ringkasan"
    ws0.append(["Simulator Node Jaringan & Evaluasi Algoritma Optimasi SDN"])
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.append([f"Diekspor: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    ws0.append([])
    ws0.append(["Bagian", "Jumlah data"])
    ws0.append(["Hasil Optimasi Routing", len(routing.get("results") or [])])
    ws0.append(["Hasil Load Balancing", len(lb.get("results") or [])])
    if routing.get("results"):
        ws0.append(["Source / Target (routing)", f"{routing.get('source')} -> {routing.get('target')}"])
        ws0.append(["Kongesti simulasi", routing.get("congestion")])
        ws0.append(["Seed acak (routing)", routing.get("seed")])
    if lb.get("results"):
        ws0.append(["Seed acak (load balancing)", lb.get("seed")])
    _autosize_columns(ws0)

    # --- Kondisi simulasi / parameter input ---
    ws_cond = wb.create_sheet("Kondisi Simulasi (Input)")
    row = 1

    def cond_section(title):
        nonlocal row
        ws_cond.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
        row += 1

    def cond_row(label, value):
        nonlocal row
        ws_cond.cell(row=row, column=1, value=label)
        ws_cond.cell(row=row, column=2, value=value)
        row += 1

    if routing.get("results"):
        topo_meta = routing.get("topology_meta") or {}
        gen_form = topo_meta.get("generation_form") or {}
        cond_section("Topologi Routing")
        cond_row("Mode topologi", topo_meta.get("mode"))
        cond_row("Jumlah node", topo_meta.get("num_nodes"))
        cond_row("Jumlah link", topo_meta.get("num_edges"))
        area = topo_meta.get("area") or {}
        if area:
            cond_row("Area (lebar x panjang, m)", f"{area.get('width')} x {area.get('height')}")
        if gen_form:
            cond_row("Parameter bangkitkan topologi -- jumlah node", gen_form.get("num_nodes"))
            cond_row("Parameter bangkitkan topologi -- tetangga terdekat (k)", gen_form.get("k_neighbors"))
            cond_row("Parameter bangkitkan topologi -- lebar wilayah (m)", gen_form.get("area_width"))
            cond_row("Parameter bangkitkan topologi -- panjang wilayah (m)", gen_form.get("area_height"))
            cond_row("Parameter bangkitkan topologi -- seed", gen_form.get("seed"))
        cond_row("Source", routing.get("source"))
        cond_row("Target", routing.get("target"))
        cond_row("Kongesti simulasi (0-1)", routing.get("congestion"))
        cond_row("Seed acak evaluasi", routing.get("seed"))
        cond_row("Algoritma dijalankan", ", ".join(
            r.get("label", r.get("algorithm")) for r in routing["results"]
        ))
        row += 1

    if lb.get("results"):
        cluster_meta = lb.get("cluster_meta") or {}
        lb_gen_form = cluster_meta.get("generation_form") or {}
        cond_section("Cluster Load Balancing")
        cond_row("Jumlah server", cluster_meta.get("num_servers"))
        cond_row("Jumlah klien", cluster_meta.get("num_clients"))
        cond_row("Request per klien", cluster_meta.get("requests_per_client"))
        lb_area = cluster_meta.get("area") or {}
        if lb_area:
            cond_row("Area (lebar x panjang, m)", f"{lb_area.get('width')} x {lb_area.get('height')}")
        if lb_gen_form:
            cond_row("Parameter bangkitkan cluster -- jumlah server", lb_gen_form.get("num_servers"))
            cond_row("Parameter bangkitkan cluster -- jumlah klien", lb_gen_form.get("num_clients"))
            cond_row("Parameter bangkitkan cluster -- request/klien", lb_gen_form.get("requests_per_client"))
            cond_row("Parameter bangkitkan cluster -- kapasitas min", lb_gen_form.get("capacity_min"))
            cond_row("Parameter bangkitkan cluster -- kapasitas maks", lb_gen_form.get("capacity_max"))
            cond_row("Parameter bangkitkan cluster -- demand min", lb_gen_form.get("demand_min"))
            cond_row("Parameter bangkitkan cluster -- demand maks", lb_gen_form.get("demand_max"))
            cond_row("Parameter bangkitkan cluster -- lebar area (m)", lb_gen_form.get("area_width"))
            cond_row("Parameter bangkitkan cluster -- panjang area (m)", lb_gen_form.get("area_height"))
            cond_row("Parameter bangkitkan cluster -- seed", lb_gen_form.get("seed"))
        cond_row("Seed acak evaluasi", lb.get("seed"))
        cond_row("Algoritma dijalankan", ", ".join(
            r.get("label", r.get("algorithm")) for r in lb["results"]
        ))

    for col_letter, width in (("A", 46), ("B", 40)):
        ws_cond.column_dimensions[col_letter].width = width

    # --- Parameter tiap algoritma (flattened) ---
    ws_params = wb.create_sheet("Parameter Algoritma")
    _write_header(ws_params, ["Konteks", "Algoritma", "Parameter", "Nilai"])
    for r in routing.get("results") or []:
        params = r.get("params") or {}
        label = r.get("label", r.get("algorithm"))
        if not params:
            ws_params.append(["Routing", label, "(tidak ada parameter)", ""])
        for k, v in params.items():
            ws_params.append(["Routing", label, k, v])
    for r in lb.get("results") or []:
        params = r.get("params") or {}
        label = r.get("label", r.get("algorithm"))
        if not params:
            ws_params.append(["Load Balancing", label, "(tidak ada parameter)", ""])
        for k, v in params.items():
            ws_params.append(["Load Balancing", label, k, v])
    _autosize_columns(ws_params)

    if routing.get("results"):
        ws = wb.create_sheet("Hasil Routing")
        _write_header(ws, [
            "Algoritma", "Source", "Target", "Kongesti", "Seed", "Tipe",
            "Jalur / Struktur", "Hop", "Delay (ms)", "Throughput (Mbps)",
            "Packet Loss (%)", "Skor QoS", "Kelas QoS", "Runtime (ms)",
            "Cost", "Parameter", "Meta", "Error",
        ])
        for r in routing["results"]:
            m = r.get("metrics") or {}
            qos = r.get("qos") or {}
            if r.get("type") == "tree":
                path_str = ", ".join("-".join(e) for e in (m.get("edges") or []))
            else:
                path_str = " -> ".join(r.get("path") or [])
            ws.append([
                r.get("label", r.get("algorithm")),
                routing.get("source"), routing.get("target"),
                routing.get("congestion"), routing.get("seed"),
                r.get("type"), path_str,
                m.get("hop_count"), m.get("delay_ms"), m.get("throughput_mbps"),
                m.get("packet_loss_pct"), qos.get("score"), qos.get("class"),
                r.get("runtime_ms"), r.get("cost"),
                json.dumps(r.get("params") or {}, ensure_ascii=False),
                json.dumps(r.get("meta") or {}, ensure_ascii=False),
                r.get("error", ""),
            ])
        _autosize_columns(ws)

    if routing.get("history"):
        ws = wb.create_sheet("Riwayat Perbandingan")
        _write_header(ws, [
            "Algoritma", "Waktu Run Sebelumnya", "Parameter Berubah",
            "Delta Delay (ms)", "Delta Hop", "Delta Throughput (Mbps)",
            "Delta Packet Loss (%)", "Delta Skor QoS",
        ])
        for h in routing["history"]:
            ws.append([
                h.get("label"), h.get("prev_ts"), h.get("changed_params"),
                h.get("delta_delay"), h.get("delta_hop"),
                h.get("delta_throughput"), h.get("delta_loss"), h.get("delta_qos"),
            ])
        _autosize_columns(ws)

    if (routing.get("topology") or {}).get("nodes"):
        topo = routing["topology"]
        ws = wb.create_sheet("Topologi (Node)")
        _write_header(ws, ["Node ID", "X", "Y"])
        for n in topo.get("nodes", []):
            ws.append([n.get("id"), n.get("x"), n.get("y")])
        _autosize_columns(ws)

        ws = wb.create_sheet("Topologi (Link)")
        _write_header(ws, ["Source", "Target", "Delay (ms)", "Bandwidth (Mbps)", "Loss"])
        for e in topo.get("edges", []):
            ws.append([e.get("source"), e.get("target"), e.get("delay"), e.get("bandwidth"), e.get("loss")])
        _autosize_columns(ws)

    if lb.get("results"):
        ws = wb.create_sheet("Hasil Load Balancing")
        _write_header(ws, [
            "Algoritma", "Seed", "Beban Maks (%)", "Rata2 Response Time (ms)",
            "Fairness (Jain)", "Dilayani", "Total", "Loss (%)",
            "Skor QoS", "Kelas QoS", "Runtime (ms)", "Cost", "Parameter", "Meta", "Error",
        ])
        for r in lb["results"]:
            m = r.get("metrics") or {}
            qos = r.get("qos") or {}
            ws.append([
                r.get("label", r.get("algorithm")), lb.get("seed"),
                m.get("max_utilization_pct"), m.get("avg_response_ms"),
                m.get("fairness_index"), m.get("served"), m.get("total"),
                m.get("dropped_pct"), qos.get("score"), qos.get("class"),
                r.get("runtime_ms"), r.get("cost"),
                json.dumps(r.get("params") or {}, ensure_ascii=False),
                json.dumps(r.get("meta") or {}, ensure_ascii=False),
                r.get("error", ""),
            ])
        _autosize_columns(ws)

    if (lb.get("cluster") or {}).get("servers"):
        c = lb["cluster"]
        ws = wb.create_sheet("Cluster (Server)")
        _write_header(ws, ["Server ID", "Kapasitas", "X", "Y"])
        for s in c.get("servers", []):
            ws.append([s.get("id"), s.get("capacity"), s.get("x"), s.get("y")])
        _autosize_columns(ws)

        ws = wb.create_sheet("Cluster (Klien)")
        _write_header(ws, ["Client ID", "X", "Y"])
        for cl in c.get("clients", []):
            ws.append([cl.get("id"), cl.get("x"), cl.get("y")])
        _autosize_columns(ws)

    # --- Log klien & server yang melayani (per algoritma load balancing) ---
    has_any_log = any((r.get("client_log") or []) for r in (lb.get("results") or []))
    if has_any_log:
        ws = wb.create_sheet("Log Klien & Server")
        _write_header(ws, [
            "Algoritma", "Client ID", "Total Request Klien", "Total Beban Klien",
            "Request Ditolak", "Server ID Dilayani", "Request ke Server", "Beban ke Server",
        ])
        for r in lb["results"]:
            label = r.get("label", r.get("algorithm"))
            for c_entry in r.get("client_log") or []:
                servers = c_entry.get("servers") or []
                if not servers:
                    ws.append([
                        label, c_entry.get("client_id"),
                        c_entry.get("total_requests"), c_entry.get("total_demand"),
                        c_entry.get("dropped_requests"),
                        "(tidak ada / semua ditolak)", 0, 0,
                    ])
                else:
                    for s_entry in servers:
                        ws.append([
                            label, c_entry.get("client_id"),
                            c_entry.get("total_requests"), c_entry.get("total_demand"),
                            c_entry.get("dropped_requests"),
                            s_entry.get("server_id"), s_entry.get("requests"), s_entry.get("demand"),
                        ])
        _autosize_columns(ws)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"sdn_simulator_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


if __name__ == "__main__":
    # Debug mode Flask (werkzeug) menampilkan modal/overlay traceback di
    # browser saat terjadi error, serta banner peringatan di awal saat
    # dijalankan. Default sekarang OFF; aktifkan dengan set env var
    # FLASK_DEBUG=1 bila butuh auto-reload/traceback interaktif saat develop.
    debug_mode = os.environ.get("FLASK_DEBUG", "0") == "1"
    app.run(debug=debug_mode, host="127.0.0.1", port=5000)
